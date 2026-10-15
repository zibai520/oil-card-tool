import streamlit as st
import pdfplumber
import pandas as pd
import re
from io import BytesIO
from openpyxl import Workbook

st.set_page_config(page_title="对账单统计工具", layout="wide")

# ====================== 功能1：各卡号交易小计汇总统计（原oil_pdf1.py） ======================
def task_card_subtotal(pdf_file):
    def safe_float(s):
        if not s:
            return 0.0
        clean_str = re.sub(r"[^\d\.\-]", "", str(s))
        try:
            return float(clean_str)
        except:
            return 0.0

    all_result = []
    with pdfplumber.open(pdf_file) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if not table:
                continue
            for row in table:
                if len(row) < 7:
                    continue
                cell0 = str(row[0]).strip()
                if "小计" in cell0:
                    card_no = str(row[2]).strip()
                    if not card_no:
                        continue
                    qty_str = str(row[6]).strip()
                    qty = safe_float(qty_str)
                    price = ""
                    all_result.append({"卡号": card_no, "加油数量": qty, "单价": price})

    df = pd.DataFrame(all_result)
    df.insert(0, "序号", range(1, len(df)+1))
    df["金额"] = df.apply(
        lambda x: round(x["加油数量"] * x["单价"], 2) if x["单价"] != "" else "",
        axis=1
    )
    total_qty = df["加油数量"].sum()
    card_count = len(df)
    total_row = pd.DataFrame([{
        "序号": "总计",
        "卡号": f"共{card_count}张卡",
        "加油数量": total_qty,
        "单价": "",
        "金额": ""
    }])
    df_final = pd.concat([df, total_row], ignore_index=True)

    # 生成Excel字节流
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_final.to_excel(writer, sheet_name="各卡号交易小计汇总", index=False)
        ws = writer.book["各卡号交易小计汇总"]
        last_data_row = len(df) + 1
        total_row_num = len(df) + 2
        for r in range(2, last_data_row + 1):
            ws.cell(row=r, column=5, value=f"=ROUND(C{r}*D{r},2)")
        ws.cell(row=total_row_num, column=3, value=f"=SUM(C2:C{last_data_row})")
        ws.cell(row=total_row_num, column=5, value=f"=ROUND(SUM(E2:E{last_data_row}),2)")
    output.seek(0)
    return df_final, output, "各卡号交易小计汇总统计.xlsx"


# ====================== 功能2：各卡号交易明细汇总统计（原oil_pdf2.py） ======================
def task_card_detail(pdf_file):
    def get_19digit_card(text):
        if text is None:
            return ""
        nums = re.findall(r"\d{19}", str(text))
        if nums:
            return nums[0]
        return ""

    def safe_parse_float(text):
        if text is None:
            return None
        s = str(text).strip()
        match = re.search(r"\d+\.\d+|\d+", s)
        if match:
            try:
                return float(match.group())
            except:
                return None
        return None

    def extract_pdf_data(pdf_f):
        all_rows = []
        main_cards = []
        with pdfplumber.open(pdf_f) as pdf:
            for page in pdf.pages:
                table = page.extract_table()
                if not table:
                    continue
                for row in table:
                    if len(row) < 3:
                        continue
                    col0 = str(row[0]).strip() if row[0] is not None else ""
                    col2 = str(row[2]).strip() if row[2] is not None else ""
                    col6 = row[6] if (len(row)>=7) else ""
                    col7 = row[7] if (len(row)>=8) else ""
                    if "小计" in col0:
                        continue
                    if "客户卡号" in col0:
                        mc = get_19digit_card(col0)
                        if mc:
                            main_cards.append(mc)
                        continue
                    if ("客户名称" in col0 or "开户省份" in col0
                            or "起止时间" in col0 or "主/副卡" in col0):
                        continue
                    if col2 == "交易时间":
                        continue
                    if "石化" in col0 and get_19digit_card(col0) == "":
                        continue
                    card = get_19digit_card(col0)
                    qty = safe_parse_float(col6)
                    price = safe_parse_float(col7)
                    all_rows.append({
                        "卡号": card,
                        "交易时间": col2,
                        "数量": qty,
                        "单价": price
                    })
        return all_rows, main_cards

    raw_list, main_cards = extract_pdf_data(pdf_file)
    df = pd.DataFrame(raw_list)
    df["卡号"] = df["卡号"].replace("", pd.NA)
    df["卡号"] = df["卡号"].ffill()
    df["单价"] = df["单价"].fillna(pd.NA)
    df["单价"] = df["单价"].ffill()
    df["数量"] = df["数量"].fillna(0.0)
    df = df.dropna(subset=["卡号"])
    all_cards = list(dict.fromkeys(df["卡号"].tolist()))
    if main_cards:
        all_cards = [c for c in all_cards if c not in main_cards]
    tx = df[(df["数量"] > 0) & (df["交易时间"] != "")].copy()

    wb = Workbook()
    ws = wb.active
    ws.append(["卡号", "交易时间", "数量", "单价", "金额"])
    curr_row = 2
    start_row = curr_row
    display = []
    subtotal_qty_cell_refs = []
    subtotal_amt_cell_refs = []
    card_subtotal_qty_values = []
    card_subtotal_amt_values = []

    for card in all_cards:
        g = tx[tx["卡号"] == card]
        total_qty_card = round(g["数量"].sum(), 2)
        total_amt_card = round((g["数量"] * g["单价"]).sum(), 2)
        if len(g) == 0:
            ws.append([card, "无交易", 0, 0, 0])
            display.append([card, "无交易", 0, 0, 0])
            curr_row += 1
            ws.append(["小计", "共0笔", 0, "", 0])
            subtotal_qty_cell_refs.append(f"C{curr_row}")
            subtotal_amt_cell_refs.append(f"E{curr_row}")
            card_subtotal_qty_values.append(0)
            card_subtotal_amt_values.append(0)
            display.append(["小计", "共0笔", 0, "", 0])
            curr_row += 1
        else:
            for _, r in g.iterrows():
                p = r["单价"] if pd.notna(r["单价"]) else 0.0
                q = r["数量"]
                ws.append([
                    r["卡号"],
                    r["交易时间"],
                    round(q,2),
                    round(p,2),
                    f"=ROUND(C{curr_row}*D{curr_row},2)"
                ])
                calc_amt = round(q * p, 2)
                display.append([r["卡号"], r["交易时间"], round(q,2), round(p,2), calc_amt])
                curr_row += 1
            qty_formula = f"=ROUND(SUM(C{start_row}:C{curr_row-1}),2)"
            amt_formula = f"=ROUND(SUM(E{start_row}:E{curr_row-1}),2)"
            ws.append(["小计", f"共{len(g)}笔", qty_formula, "", amt_formula])
            subtotal_qty_cell_refs.append(f"C{curr_row}")
            subtotal_amt_cell_refs.append(f"E{curr_row}")
            card_subtotal_qty_values.append(total_qty_card)
            card_subtotal_amt_values.append(total_amt_card)
            display.append(["小计", f"共{len(g)}笔", total_qty_card, "", total_amt_card])
            curr_row += 1
        start_row = curr_row

    total_qty_formula = f"=ROUND(SUM({','.join(subtotal_qty_cell_refs)}),2)"
    total_amt_formula = f"=ROUND(SUM({','.join(subtotal_amt_cell_refs)}),2)"
    total_trade_count = len(tx)
    ws.append(["总计", f"共{total_trade_count}笔", total_qty_formula, "", total_amt_formula])
    sum_total_qty = round(sum(card_subtotal_qty_values),2)
    sum_total_amt = round(sum(card_subtotal_amt_values),2)
    display.append(["总计", f"共{total_trade_count}笔", sum_total_qty, "", sum_total_amt])

    df_display = pd.DataFrame(display, columns=["卡号", "交易时间", "数量", "单价", "金额"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return df_display, buf, "各卡号交易明细汇总统计.xlsx"


# ===================== 页面主体UI =====================
st.title("📋 对账单统计工具")

mode = st.selectbox(
    "请选择统计功能",
    [
        "各卡号交易小计汇总统计",
        "各卡号交易明细汇总统计"
    ]
)

if mode == "各卡号交易小计汇总统计":
    st.subheader("🔹 各卡号交易小计汇总统计")
    upload_pdf = st.file_uploader("上传PDF对账单", type="pdf")
    if upload_pdf is not None:
        with st.spinner("正在解析PDF并汇总..."):
            res_df, excel_bytes, filename = task_card_subtotal(upload_pdf)
        st.dataframe(res_df, use_container_width=True, hide_index=True)
        st.download_button("下载Excel", data=excel_bytes, file_name=filename,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.warning("⚠️  若 PDF 报表排版发生变动，有可能造成识别异常，每次统计完成后，请抽样核对2~3 个卡号的明细数据，同时核对总计数量，确保结果无误。")

elif mode == "各卡号交易明细汇总统计":
    st.subheader("🔹 各卡号交易明细汇总统计")
    upload_pdf = st.file_uploader("上传PDF对账单", type="pdf")
    if upload_pdf is not None:
        with st.spinner("正在解析PDF并汇总..."):
            res_df, excel_bytes, filename = task_card_detail(upload_pdf)
        st.dataframe(res_df, use_container_width=True, hide_index=True)
        st.download_button("下载Excel", data=excel_bytes, file_name=filename,
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.warning("⚠️ 若 PDF 报表排版发生变动，有可能造成识别异常，每次统计完成后，请抽样核对2~3 个卡号的明细数据，同时核对总计数量，确保结果无误。")

